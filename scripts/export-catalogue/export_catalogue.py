"""Export EN LECTURE SEULE du catalogue Tennis String Advisor vers Excel.

Une commande, depuis la racine du dépôt :
    npm run export:catalogue
(équivaut à : python scripts/export-catalogue/export_catalogue.py --repo .)
Sortie : exports/catalogue_TSA_<date>.xlsx (dossier non versionné).
Prérequis : Python 3 + openpyxl ; Node >= 22.6 (import natif des fichiers .ts).

Étapes : (1) node dump-catalogue.mjs lit Supabase (PostgREST, clé anon, GET seulement,
count exact vérifié) et charge les modules TS servis par le site FR ; (2) ce script écrit
le classeur. Aucune valeur n'est comblée : un champ absent reste une cellule vide.
"""
import argparse, datetime as dt, json, os, re, subprocess, sys, tempfile
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
URL_RE = re.compile(r"https?://[^\s,)»]+")
BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDE5F0")
LINK = Font(color="0563C1", underline="single")


def cell_value(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return v
    if isinstance(v, list):
        return " | ".join("" if x is None else str(x) for x in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return v


def flatten(obj, prefix=""):
    """Aplatit les objets imbriqués (price.europe…) sans rien inventer."""
    out = {}
    for k, v in obj.items():
        key = f"{prefix}{k}"
        if isinstance(v, dict):
            out.update(flatten(v, key + "."))
        else:
            out[key] = v
    return out


def write_table(ws, columns, rows, url_cols=()):
    ws.append(columns)
    for c in ws[1]:
        c.font = BOLD
        c.fill = HEAD_FILL
    for r in rows:
        ws.append([cell_value(r.get(c)) for c in columns])
    for ci, col in enumerate(columns, 1):
        if col in url_cols:
            for ri in range(2, ws.max_row + 1):
                cell = ws.cell(ri, ci)
                if isinstance(cell.value, str) and cell.value.startswith("http"):
                    cell.hyperlink = cell.value
                    cell.font = LINK
        # largeur raisonnable : bornée entre 8 et 60
        lens = [len(str(col))] + [len(str(ws.cell(ri, ci).value or "").split("\n")[0]) for ri in range(2, min(ws.max_row, 400) + 1)]
        ws.column_dimensions[get_column_letter(ci)].width = max(8, min(60, max(lens) + 2))
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def ts_rows(block):
    rows, cols = [], []
    max_urls = 0
    for obj in block["rows"]:
        flat = flatten(obj)
        for k in flat:
            if k not in cols:
                cols.append(k)
        cm = block["comments"][obj["id"]]
        urls = URL_RE.findall(cm["lead"] + "\n" + cm["inner"])
        flat["ts_commentaire_avant_objet"] = cm["lead"] or None
        flat["ts_commentaires_dans_objet"] = cm["inner"] or None
        flat["ts_section_amont (dernier bandeau ====)"] = cm["section"] or None
        flat["ts_mention_TWU"] = "oui" if "TWU" in (cm["lead"] + cm["inner"]) else None
        for i, u in enumerate(urls, 1):
            flat[f"ts_url_{i}"] = u
        max_urls = max(max_urls, len(urls))
        rows.append(flat)
    prov = ["ts_commentaire_avant_objet", "ts_commentaires_dans_objet",
            "ts_section_amont (dernier bandeau ====)", "ts_mention_TWU"] + [f"ts_url_{i}" for i in range(1, max_urls + 1)]
    return cols + prov, rows, [f"ts_url_{i}" for i in range(1, max_urls + 1)]


# Correspondance Supabase -> TS (nom réel des deux côtés, aucune transformation de valeur)
MAP = {
    "racquets": [("brand", "brand"), ("model", "model"), ("variant", "variant"), ("stiffness", "stiffness"),
                 ("weight", "weight"), ("head_size", "headSize"), ("string_pattern", "stringPattern"),
                 ("category", "category"), ("player_level", "playerLevel"), ("description", "description"),
                 ("pro_usage", "proUsage"), ("price_eur", "price.europe"), ("price_usd", "price.usa")],
    "strings": [("brand", "brand"), ("model", "model"), ("type", "type"), ("gauges", "gauges"),
                ("stiffness", "stiffness"), ("performance", "performance"), ("control", "control"),
                ("comfort", "comfort"), ("durability", "durability"), ("spin", "spin"), ("power", "power"),
                ("tension_min", "recommendedTension.min"), ("tension_max", "recommendedTension.max"),
                ("price_eur", "price.europe"), ("price_usd", "price.usa"), ("description", "description"),
                ("pro_usage", "proUsage"), ("color", "color")],
}


def same(a, b):
    if a in (None, "") and b in (None, ""):
        return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return abs(float(a) - float(b)) < 1e-9
    if isinstance(a, list) and isinstance(b, list):
        return [str(x) for x in a] == [str(x) for x in b]
    return a == b


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", default=os.getcwd())
    ap.add_argument("--out", default=None)
    a = ap.parse_args()
    repo = os.path.abspath(a.repo)
    today = dt.date.today().isoformat()
    out = a.out or os.path.join(repo, "exports", f"catalogue_TSA_{today}.xlsx")
    os.makedirs(os.path.dirname(out), exist_ok=True)

    tmp = os.path.join(tempfile.gettempdir(), f"tsa_catalogue_dump_{os.getpid()}.json")
    subprocess.run(["node", "--no-warnings", os.path.join(HERE, "dump-catalogue.mjs"), repo, tmp], check=True)
    with open(tmp, encoding="utf-8") as f:
        d = json.load(f)
    os.remove(tmp)

    wb = Workbook()
    readme = wb.active
    readme.title = "Lisez-moi"
    sheets = {}

    for kind, name in (("racquets", "Supabase_Raquettes"), ("strings", "Supabase_Cordages")):
        t = d["supabase"][kind]
        ws = wb.create_sheet(name)
        write_table(ws, t["columns"], t["rows"])
        sheets[name] = (len(t["rows"]), t["count_exact"])

    ts_flat = {}
    for kind, name in (("racquets", "TS_Raquettes"), ("strings", "TS_Cordages")):
        cols, rows, urlc = ts_rows(d["ts"][kind])
        ws = wb.create_sheet(name)
        write_table(ws, cols, rows, url_cols=urlc)
        sheets[name] = (len(rows), len(d["ts"][kind]["rows"]))
        ts_flat[kind] = {r["id"]: r for r in rows}

    # --- Écarts
    ws = wb.create_sheet("Ecarts")
    det = wb.create_sheet("Ecarts_detail")
    summary, detail, stats = [], [], {}
    removed = set(d["ts"]["strings"]["removed_ids"])
    aliases = d["ts"]["strings"]["legacy_aliases"]
    for kind, label in (("racquets", "raquette"), ("strings", "cordage")):
        sb = {r["id"]: r for r in d["supabase"][kind]["rows"]}
        tsr = ts_flat[kind]
        only_sb = sorted(set(sb) - set(tsr))
        only_ts = sorted(set(tsr) - set(sb))
        common = sorted(set(sb) & set(tsr))
        n_conf_ids = n_rcs = n_conf_nodesc = 0
        for i in only_sb:
            note = None
            if kind == "strings" and i in removed:
                note = "id retiré du TS (REMOVED_STRING_IDS)"
            elif kind == "strings" and i in aliases:
                note = f"alias legacy dans le TS -> {aliases[i]}"
            summary.append({"type": label, "id": i, "presence": "Supabase seulement", "note": note})
        for i in only_ts:
            summary.append({"type": label, "id": i, "presence": "TS seulement"})
        for i in common:
            conf = [(s, t) for s, t in MAP[kind] if not same(sb[i].get(s), tsr[i].get(t))]
            if conf:
                n_conf_ids += 1
            if any(s != "description" for s, _ in conf):
                n_conf_nodesc += 1
            if any(s == "stiffness" for s, _ in conf):
                n_rcs += 1
            summary.append({"type": label, "id": i, "presence": "commun", "nb_champs_en_conflit": len(conf),
                            "champs_en_conflit": ", ".join(s for s, _ in conf) or None,
                            "nb_conflits_hors_description": sum(1 for s, _ in conf if s != "description"),
                            "rigidite_differe (entree RCS)": "oui" if any(s == "stiffness" for s, _ in conf) else None})
            for s, t in conf:
                detail.append({"type": label, "id": i, "champ_supabase": s, "champ_ts": t,
                               "valeur_supabase": sb[i].get(s), "valeur_ts": tsr[i].get(t)})
        stats[kind] = dict(sb=len(sb), ts=len(tsr), only_sb=len(only_sb), only_ts=len(only_ts), common=len(common),
                           conf_ids=n_conf_ids, conf_ids_hors_desc=n_conf_nodesc, conf_fields=sum(1 for x in detail if x["type"] == label), rcs=n_rcs)
    write_table(ws, ["type", "id", "presence", "nb_champs_en_conflit", "champs_en_conflit", "nb_conflits_hors_description",
                     "rigidite_differe (entree RCS)", "note"], summary)
    write_table(det, ["type", "id", "champ_supabase", "champ_ts", "valeur_supabase", "valeur_ts"], detail)

    # --- Lisez-moi
    r, s = d["ts"]["racquets"], d["ts"]["strings"]
    ref = d["supabase"]["url"].split("//")[1].split(".")[0]
    lines = [
        ("Catalogue Tennis String Advisor — extraction du " + today, None),
        ("", None),
        ("AVERTISSEMENT — DIVERGENCE", None),
        ("1. Les fichiers TypeScript (onglets TS_*) sont le catalogue de référence du site FR et, depuis le chantier C3 (09/10/2026), des pages EN. Les tables Supabase (onglets Supabase_*) ne sont plus lues par le site une fois C3 déployé : elles sont figées.", None),
        (f"2. Supabase n'a pas été mis à jour depuis les corrections de sept. 2026 : {stats['racquets']['common']} raquettes communes seulement sur {stats['racquets']['sb']} (Supabase) / {stats['racquets']['ts']} (TS) ; cordages : {stats['strings']['conf_ids_hors_desc']} des {stats['strings']['common']} ids communs ont au moins une valeur en conflit hors description, dont {stats['strings']['rcs']} sur la rigidité (entrée du RCS).", None),
        ("3. Supabase ne possède AUCUNE colonne de source, d'URL ou de provenance. Dans le TS, la provenance n'existe que sous forme de commentaires (colonnes ts_*), présente sur une minorité de produits.", None),
        ("", None),
        ("ORIGINE DES ONGLETS", None),
        (f"Supabase_Raquettes : projet {ref}, table public.racquets, API REST (clé anon, lecture seule)", f"{sheets['Supabase_Raquettes'][0]} lignes exportées / count(*) exact = {sheets['Supabase_Raquettes'][1]}"),
        (f"Supabase_Cordages : projet {ref}, table public.strings, API REST (clé anon, lecture seule)", f"{sheets['Supabase_Cordages'][0]} lignes exportées / count(*) exact = {sheets['Supabase_Cordages'][1]}"),
        (f"TS_Raquettes : {r['file']}, export {r['export']} — branche {r['branch']} @ {r['head']}, dernier commit du fichier {r['last_commit']}{' (MODIFIÉ localement)' if r['dirty'] else ''}", f"{sheets['TS_Raquettes'][0]} lignes / {sheets['TS_Raquettes'][1]} objets exportés"),
        (f"TS_Cordages : {s['file']}, export {s['export']} — branche {s['branch']} @ {s['head']}, dernier commit du fichier {s['last_commit']}{' (MODIFIÉ localement)' if s['dirty'] else ''}", f"{sheets['TS_Cordages'][0]} lignes / {sheets['TS_Cordages'][1]} objets exportés"),
        ("Ecarts / Ecarts_detail : comparaison par id, champ à champ (correspondance des noms Supabase ↔ TS ; aucune valeur convertie)", None),
        ("", None),
        ("ÉCARTS — SYNTHÈSE", None),
    ]
    for kind, lab in (("racquets", "Raquettes"), ("strings", "Cordages")):
        x = stats[kind]
        lines.append((f"{lab} : Supabase seulement {x['only_sb']} · TS seulement {x['only_ts']} · communs {x['common']} (dont {x['conf_ids']} avec au moins un champ en conflit, {x['conf_ids_hors_desc']} hors description ; {x['conf_fields']} champs en conflit au total ; {x['rcs']} ids dont la rigidité diffère)", None))
    lines += [
        ("", None),
        ("CONVENTIONS", None),
        ("Cellule vide = valeur absente à la source. Aucune valeur n'a été comblée, convertie ni renommée.", None),
        ("Listes (gauges, player_level, playerLevel) : éléments séparés par « | ». Objets imbriqués du TS aplatis : price.europe, price.usa, recommendedTension.min/max.", None),
        ("ts_commentaire_avant_objet / ts_commentaires_dans_objet : commentaires du code source attachés au produit, verbatim. ts_section_amont : dernier bandeau « ==== » au-dessus du produit (s'applique à un groupe ; rattachement par position).", None),
        ("ts_url_N : URLs citées dans ces commentaires (cliquables). ts_mention_TWU : « oui » si le commentaire cite une mesure TWU.", None),
        ("Champs TS versatility et innovation : absents du schéma Supabase, donc hors comparaison. created_at : Supabase seulement.", None),
        (f"Ids cordages retirés du TS (REMOVED_STRING_IDS) : {', '.join(s['removed_ids'])}", None),
        ("Alias legacy TS : " + ", ".join(f"{k} -> {v}" for k, v in s["legacy_aliases"].items()), None),
        ("", None),
        (f"Généré par export_catalogue.py le {d['generated_at']} (UTC). Relance : python export_catalogue.py --repo <dépôt>", None),
    ]
    for a_, b_ in lines:
        readme.append([a_, b_])
    for row in readme.iter_rows():
        if row[0].value and (row[0].value.isupper() or row[0].value.startswith(("AVERTISSEMENT", "ORIGINE", "ÉCARTS", "CONVENTIONS", "Catalogue"))):
            row[0].font = BOLD
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    readme["A1"].font = Font(bold=True, size=14)
    readme.column_dimensions["A"].width = 130
    readme.column_dimensions["B"].width = 50

    for ws in wb.worksheets[1:]:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, str) and "\n" in c.value:
                    c.alignment = Alignment(wrap_text=False, vertical="top")
    wb.save(out)
    print("Écrit :", out)
    print(json.dumps({"sheets": sheets, "ecarts": stats}, ensure_ascii=False))


if __name__ == "__main__":
    main()
